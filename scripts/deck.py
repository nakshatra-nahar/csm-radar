"""The Radar deck, docs/deck/radar.pptx, built from docs/img and styled like the app:

    uv run --with python-pptx --with pillow python scripts/deck.py [--out docs/deck/radar.pptx]

python-pptx and Pillow are not project dependencies; they are only needed here. After saving,
the deck is read back: every shape inside the 16:9 slide, every text box tall enough for its
text (estimated from average glyph widths), notes on every slide. A problem exits non-zero.
Figures on the slides come from eval/*.json and docs/measurements.md.
"""

import argparse
import math
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "docs" / "img"
W, H = Inches(13.333), Inches(7.5)
BG, INK, MUTED = RGBColor(0xF6, 0xF3, 0xEC), RGBColor(0x1F, 0x1F, 0x1F), RGBColor(0x5A, 0x5A, 0x5A)
TEAL, RED, CARD = RGBColor(0x0F, 0x5C, 0x5C), RGBColor(0xB3, 0x26, 0x1E), RGBColor(0xFF, 0xFD, 0xF8)
LINE = RGBColor(0xD9, 0xD4, 0xC7)
SERIF, SANS = "Georgia", "Arial"
LEFT = Inches(0.75)
CHAR_EM = {SERIF: 0.56, SANS: 0.53}  # generous average glyph widths, for the fit check


def deck() -> Presentation:
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    # the default template says type="screen4x3"; without it PowerPoint reads the size as custom
    prs.part._element.find(qn("p:sldSz")).attrib.pop("type", None)
    return prs


def slide(prs, title: str | None, notes: str):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    rule = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.4), 0, Inches(0.05), H)
    rule.fill.solid()
    rule.fill.fore_color.rgb = RED
    rule.line.fill.background()
    if title:
        text(s, LEFT, Inches(0.45), W - LEFT - Inches(0.6), Inches(0.9),
             [(title, 28, SERIF, INK, False)])
    s.notes_slide.notes_text_frame.text = notes
    return s


def text(s, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT, fill=None):
    """paras: (text, size, font, colour, bold) or (runs, size) where runs are (text, font,
    colour, bold) for a bold lead-in."""
    box = s.shapes.add_textbox(x, y, w, h) if fill is None else s.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, x, y, w, h)
    if fill is not None:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.color.rgb = LINE
    tf = box.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = Inches(0.06)
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(para[1] * 0.5)
        runs = para[0] if isinstance(para[0], list) else [(para[0], para[2], para[3], para[4])]
        for t, font, colour, bold in runs:
            r = p.add_run()
            r.text = t
            r.font.size, r.font.name, r.font.bold = Pt(para[1]), font, bold
            r.font.color.rgb = colour
    return box


def body(t, size=15, colour=INK, bold=False):
    return (t, size, SANS, colour, bold)


def lead(head, rest, size=15):
    return ([(head, SANS, TEAL, True), (rest, SANS, INK, False)], size)


def picture(s, name, x, y, w, h, crop=None):
    """Fit docs/img/<name> (or its crop box, in image pixels) inside the box, keeping its
    aspect ratio; returns (shape, scale in EMU per image pixel)."""
    path = IMG / name
    iw, ih = Image.open(path).size
    left, top, right, bottom = crop or (0, 0, iw, ih)
    scale = min(w / (right - left), h / (bottom - top))
    pw, ph = int((right - left) * scale), int((bottom - top) * scale)
    pic = s.shapes.add_picture(str(path), x + (w - pw) // 2, y, pw, ph)
    pic.crop_left, pic.crop_top = left / iw, top / ih
    pic.crop_right, pic.crop_bottom = 1 - right / iw, 1 - bottom / ih
    pic.line.color.rgb = LINE
    return pic, scale


def marker(s, n, x, y):
    dot = s.shapes.add_shape(MSO_SHAPE.OVAL, x, y, Inches(0.28), Inches(0.28))
    dot.fill.solid()
    dot.fill.fore_color.rgb = TEAL
    dot.line.color.rgb = CARD
    tf = dot.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size, r.font.name, r.font.bold = Pt(11), SANS, True
    r.font.color.rgb = CARD
    return dot


def needed_height(paras, tf, width) -> int:
    """Estimated height (EMU) of [(text, size_pt, font, space_after_pt)] in text frame tf at
    this box width: wraps by average glyph width, honours line breaks and the frame's insets."""
    usable = width - tf.margin_left - tf.margin_right
    total = tf.margin_top + tf.margin_bottom
    for t, size, font, after in paras:
        per_line = max(1, int(usable / (Pt(size) * CHAR_EM.get(font, 0.56))))
        lines = sum(max(1, math.ceil(len(seg) / per_line)) for seg in t.split("\v"))
        total += lines * Pt(size) * 1.2 + Pt(after)
    return int(total)


def check(path) -> tuple[list[str], list[str]]:
    """Read the saved deck back: a 16:9 slide size with no 4:3 type, every shape inside the
    slide, every text box tall enough for its text, notes on every slide. Returns (problems, a
    text dump of every slide)."""
    prs, problems, dump = Presentation(path), [], []
    size = prs.part._element.find(qn("p:sldSz"))
    if (prs.slide_width, prs.slide_height) != (W, H) or size.get("type"):
        problems.append(f"slide size {size.attrib}")
    for i, s in enumerate(prs.slides, 1):
        notes = s.notes_slide.notes_text_frame.text.strip()
        if not notes:
            problems.append(f"slide {i}: no notes")
        dump.append(f"--- slide {i}")
        for sh in s.shapes:
            box = (Emu(sh.left).inches, Emu(sh.top).inches, Emu(sh.width).inches,
                   Emu(sh.height).inches)
            if sh.left < 0 or sh.top < 0 or sh.left + sh.width > W or sh.top + sh.height > H:
                problems.append(f"slide {i}: {sh.name} off-slide {box}")
            if sh.shape_type == 13:  # picture
                dump.append(f"[picture {box[0]:.2f},{box[1]:.2f} {box[2]:.2f}x{box[3]:.2f}]")
            if not sh.has_text_frame or not sh.text_frame.text.strip():
                continue
            ps = [p for p in sh.text_frame.paragraphs if p.runs]
            paras = [(p.text, p.runs[0].font.size.pt, p.runs[0].font.name,
                      p.space_after.pt if p.space_after is not None and p is not ps[-1] else 0)
                     for p in ps]
            need = needed_height(paras, sh.text_frame, sh.width)
            if need > sh.height:
                problems.append(f"slide {i}: {sh.name} text needs {Emu(need).inches:.2f} in,"
                                f" box {box[3]:.2f} in: {paras[0][0][:40]!r}")
            dump += [f"  {t}" for t, *_ in paras]
        dump.append(f"  notes: {notes}")
    return problems, dump


def s1(prs):
    s = slide(prs, None, "Radar is what I built for the take-home. Every Monday it gives each CSM "
              "a short list of the accounts to call, the reason for each one, and a draft email. "
              "All the data is synthetic: three customer organisations, 5,000 accounts and about "
              "230,000 records across five sources, messy and out of order on purpose.")
    text(s, LEFT, Inches(2.0), Inches(11), Inches(1.6), [("Radar", 66, SERIF, INK, False)])
    text(s, LEFT, Inches(3.7), Inches(12), Inches(0.8),
         [("A Monday worklist for CSMs, built for the Fello take-home", 26, SERIF, TEAL, False)])
    text(s, LEFT, Inches(5.2), Inches(11), Inches(0.9),
         [body("Nakshatra Nahar · Staff Engineer, Data Architecture & GTM AI Platform", 16,
               MUTED),
          body("October 2026 · all data synthetic", 16, MUTED)])


def s2(prs):
    s = slide(prs, "The real problem is a Monday decision",
              "The VP asked for red, yellow and green per account, and an automatic email when an "
              "account goes red. I read the real problem as a choice each CSM makes every Monday: "
              "which ten of their accounts get a call this week. Colours don't rank a book, and an "
              "automatic email to a customer with an angry open ticket does damage, so I built a "
              "ranked list with evidence and drafts that a human approves.")
    text(s, LEFT, Inches(1.6), Inches(5.6), Inches(0.5), [body("The request", 14, MUTED, True)])
    text(s, LEFT, Inches(2.05), Inches(5.6), Inches(3.1), [
        ("“Every Monday they open a spreadsheet and basically guess who to call.”", 19,
         SERIF, INK, False),
        ("“What I want is a dashboard with a health score per account, red / yellow / "
         "green. And then automate the outreach — if an account goes red, the system should "
         "email them.”", 19, SERIF, INK, False),
        body("VP Customer Success", 14, MUTED)], fill=CARD)
    x = LEFT + Inches(6.1)
    text(s, x, Inches(1.6), Inches(5.8), Inches(0.5),
         [body("What I think the real problem is", 14, MUTED, True)])
    text(s, x, Inches(2.05), Inches(5.8), Inches(4.6), [
        lead("A capacity problem. ", "A CSM owns about 120 accounts and has room for about ten "
             "proactive calls a week."),
        lead("A ranking problem. ", "CSMs need to know who to call this week and why, with "
             "evidence they can trust. Thirty red accounts is still a guess about who comes "
             "first."),
        lead("A trust problem. ", "An automatic email to a red account can reach a customer with "
             "an angry open ticket, or a champion who left last month."),
        lead("So: ", "a ranked weekly list sized to capacity, every risk with an owner and a "
             "due date, and drafts that a human sends.")])


def s3(prs):
    s = slide(prs, "What a CSM sees on Monday",
              "This is My week for one hearthline CSM, from the full-size database. Hard facts "
              "come first as must-act cases with a due date. A check-in covers an account we "
              "know too little about that renews soon. Then the model's slots fill the list to "
              "ten, each with its top behavioural reason and a draft that waits for the CSM.")
    pic, k = picture(s, "week.png", LEFT, Inches(1.45), Inches(7.6), Inches(5.75))
    for n, (px, py) in enumerate(((62, 248), (62, 778), (1056, 318)), 1):  # week.png pixels
        marker(s, n, pic.left + int(px * k), pic.top + int(py * k))
    x = LEFT + Inches(7.9)
    text(s, x, Inches(1.45), Inches(4.2), Inches(5.75), [
        lead("1  Must act. ", "A hard fact (a cancellation, an urgent ticket stalled for three "
             "days) opens a case with a due date. It skips the model.", 14),
        lead("2  Model slots. ", "The rest of the ten, ranked by expected ARR at risk. Each row "
             "shows its biggest behavioural reason, such as “support friction (90 days)”.", 14),
        lead("3  A draft, not an email. ", "Every row has a play and a draft. Nothing goes to "
             "the customer until the CSM approves it, and the send is a dry run.", 14),
        body("The list is frozen on Monday. New must-acts are added mid-week with a NEW mark; "
             "the list is never re-sorted.", 13, MUTED)])


def s4(prs):
    s = slide(prs, "How it decides: rules first, then the model",
              "Hard facts never wait for a model: they open a case with an owner and a due date. "
              "Then up to two check-ins for accounts we can't see that renew soon. The model fills "
              "what is left, ranked by p times ARR, and a slot needs at least the tenant's mean "
              "expected ARR at risk. Capacity is a cap, so a healthy book gets a shorter list, and "
              "the floor is on the ranking key, so it never reorders a list. The 90 under a floor "
              "on p is in docs/measurements.md, from the earlier eval at commit f517b6c.")
    steps = [("1  Must-act cases, all of them",
              "Cancel announced, dunning exhausted, champion email bounced, urgent ticket stalled. "
              "Each has an SLA. More than K shows as overflow on the team page."),
             ("2  Up to two check-ins",
              "Accounts with Unknown coverage whose term ends within 90 days."),
             ("3  Model slots, up to K = 10",
              "Ranked by expected ARR at risk (p × ARR). A slot needs at least the "
              "tenant's mean expected ARR at risk.")]
    for i, (head, rest) in enumerate(steps):
        y = Inches(1.6) + i * Inches(1.5)
        text(s, LEFT, y, Inches(6.3), Inches(1.3),
             [(head, 17, SERIF, TEAL, False), body(rest, 14)], fill=CARD)
    x = LEFT + Inches(6.8)
    text(s, x, Inches(1.6), Inches(5.2), Inches(5.3), [
        body("Why the floor sits on the ranking key", 14, MUTED, True),
        body("Capacity is a cap, not a quota: a healthy book gets a shorter list.", 15),
        body("Each list is a prefix of its book's ranking. The floor can shorten a list but "
             "never reorders it or swaps in a smaller account.", 15),
        body("The cost: it catches more churned ARR but flags fewer small churners early. In "
             "tallyworks, 70 of 174 decisions were flagged ahead, against 90 under a floor on p.",
             15),
        body("Small high-risk accounts still reach the CSM through the must-act rules and the "
             "Book page.", 15)])


def s5(prs):
    s = slide(prs, "Every number traces back to the raw records",
              "On the left is the account page: the why is the model's contributions in log-odds, "
              "grouped by source, and the brief cites evidence items E1 to E5. On the right is the "
              "trace for the same score: the run, the model version, the as-of time, the source "
              "watermarks and the code hash of every feature. Recompute from raw rebuilds the "
              "features and p from stored events with that run's model, and here it matches.")
    picture(s, "account.png", LEFT, Inches(1.4), Inches(5.4), Inches(5.8),
            crop=(120, 60, 1330, 1350))
    picture(s, "trace.png", LEFT + Inches(5.7), Inches(1.4), Inches(6.3), Inches(4.1),
            crop=(120, 0, 1330, 770))
    text(s, LEFT + Inches(5.7), Inches(5.5), Inches(6.3), Inches(1.75), [
        lead("Why: ", "contributions per feature in log-odds. They add up to the score's logit, "
             "checked on every stored score.", 14),
        lead("Trace: ", "run, model version, as-of time, watermarks, a code hash per feature.",
             14),
        lead("Recompute from raw: ", "rebuilds features and p from stored events with the "
             "run's model. Result: match.", 14)])


def s6(prs):
    s = slide(prs, "Does it work? More churned ARR on the lists",
              "Recall at K is the share of the ARR that went on to churn that was on a CSM's top "
              "ten that week, from a point-in-time backtest with rolling origins. The model beats "
              "the rules by about eight and nine points, and both intervals stay above zero. "
              "Renewal-date sort does worse than random in hearthline, where its interval sits "
              "below random's, and no better than random in tallyworks, where the intervals "
              "overlap: accounts close to renewal have mostly already announced their decision "
              "and left the list. The VP's two lost renewals were on the model's lists 12 and 11 "
              "weeks ahead; renewal sort found them 2 weeks and 1 week ahead.")
    picture(s, "recall.png", LEFT, Inches(1.45), Inches(7.4), Inches(4.3))
    text(s, LEFT, Inches(5.45), Inches(7.4), Inches(1.6), [
        body("Recall@K: share of churned ARR that was on a CSM's top-10 list that week. "
             "Point-in-time backtest, rolling origins, paired cluster bootstrap over accounts "
             "(1,000 resamples). Kitewire, onboarded this year, has no model yet.", 12, MUTED)])
    x = LEFT + Inches(7.7)
    text(s, x, Inches(1.45), Inches(4.4), Inches(5.6), [
        lead("hearthline: ", "model 61.4% (52.9 to 68.0), rules 53.5%, renewal sort 14.7%. "
             "Model minus rules +7.9 points (+2.9 to +12.9).", 14),
        lead("tallyworks: ", "model 59.1% (49.3 to 67.0), rules 49.7%, renewal sort 18.4%. "
             "Model minus rules +9.4 points (+2.0 to +17.3).", 14),
        lead("Renewal sort does worse than random in hearthline and no better than random in "
             "tallyworks, ", "because accounts close to renewal have mostly announced their "
             "decision already and left the list.", 14),
        lead("The VP's two lost renewals: ", "on the model's list 12 and 11 weeks before the "
             "decision; renewal sort found them 2 weeks and 1 week before.", 14)])


def box(s, x, y, w, h, title, sub, colour=TEAL):
    return text(s, x, y, w, h, [(title, 14, SANS, colour, True), (sub, 11, SANS, INK, False)],
                anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.CENTER, fill=CARD)


def arrow(s, x, y):
    a = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, x, y, Inches(0.22), Inches(0.24))
    a.fill.solid()
    a.fill.fore_color.rgb = MUTED
    a.line.fill.background()


def s7(prs):
    s = slide(prs, "Architecture: tenants kept apart by the database",
              "Records land unchanged in raw_records, then become events with two clocks: when "
              "something happened and when we could first have known it. Every feature is read "
              "as of a time T, so training never sees the future. Identity uses exact keys only. "
              "Isolation is Postgres row-level security, forced on every tenant table, and the "
              "app's roles cannot bypass it; each tenant gets its own model.")
    stages = [("Sources", "Salesforce, Stripe, Zendesk, Gmail metadata, product"),
              ("raw_records", "immutable ledger, one row per source record"),
              ("events", "two clocks: occurred_at and available_at"),
              ("identity", "exact keys only; links are append-only"),
              ("features", "13, read as of T"),
              ("model", "logistic regression per tenant, JSON artifact"),
              ("policy", "must-acts, check-ins, model slots"),
              ("worklist", "frozen each Monday")]
    w, gap, y = Inches(1.19), Inches(0.32), Inches(1.75)
    for i, (t, sub) in enumerate(stages):
        x = LEFT + i * (w + gap)
        box(s, x, y, w, Inches(1.75), t, sub)
        if i < len(stages) - 1:
            arrow(s, x + w + Inches(0.05), y + Inches(0.75))
    cols = [("As-of everywhere", "A record is visible at T only if available_at ≤ T and "
             "occurred_at ≤ T. Training, scoring, evidence and the trace use the same rule."),
            ("Tenant isolation in Postgres", "Every tenant table has FORCE row-level security on "
             "tenant_id. The app roles are NOBYPASSRLS and own nothing; the app refuses to start "
             "with an owner DSN."),
            ("A model per tenant", "Each tenant's model is fitted on its own history. Kitewire, "
             "onboarded this year, starts cold on the rules alone.")]
    for i, (head, rest) in enumerate(cols):
        x = LEFT + i * Inches(4.0)
        text(s, x, Inches(4.0), Inches(3.7), Inches(2.9),
             [(head, 16, SERIF, RED if i == 1 else TEAL, False), body(rest, 14)])


def s8(prs):
    s = slide(prs, "An unattended agent, and a human on every email",
              "Nobody presses a button for any of this. A payment failure or a stalled ticket "
              "queues a rule check in the same transaction that stores it, a daily sweep rescores "
              "every account, Monday freezes the lists, and an hourly check escalates missed due "
              "dates once. The LLM only writes the wording: a deterministic validator checks it, "
              "a template replaces it if it fails twice, and nothing reaches a customer until a "
              "CSM approves the draft.")
    text(s, LEFT, Inches(1.55), Inches(5.8), Inches(0.5),
         [body("What runs unattended", 14, MUTED, True)])
    text(s, LEFT, Inches(2.0), Inches(5.8), Inches(5.0), [
        lead("Event-driven rules. ", "A payment failure, subscription change, ticket change or "
             "bounce queues a rule check in the same transaction. A must-act case opens at the "
             "next worker tick.", 14),
        lead("Daily sweep ", "at 06:00 tenant time: score every account, recheck every rule.",
             14),
        lead("Monday freeze: ", "lock each list, open cases, write briefs and digests.", 14),
        lead("SLA escalation: ", "an hourly check sends each overdue case to the manager, once.",
             14),
        lead("Circuit breaker: ", "a burst of new cases is held behind one incident, so a late "
             "batch can't become hundreds of alerts.", 14)])
    x = LEFT + Inches(6.3)
    text(s, x, Inches(1.55), Inches(5.6), Inches(0.5),
         [body("Guardrails on the LLM", 14, MUTED, True)])
    text(s, x, Inches(2.0), Inches(5.6), Inches(3.1), [
        lead("Words only. ", "The LLM never scores, ranks, picks the recipient or picks the "
             "play. Code does.", 14),
        lead("A deterministic validator. ", "Every number must appear in the cited evidence; "
             "no banned words (churn, risk, discount); no URLs or addresses in the email.", 14),
        lead("Template fallback. ", "One retry with the errors, then a fixed template. No key "
             "or no budget means the template.", 14),
        lead("Drafts only. ", "Every customer email waits in the outbox until a CSM approves "
             "it. The transport is a dry run and the deploy has no SMTP credentials.", 14)],
         fill=CARD)


def s9(prs):
    s = slide(prs, "What breaks first at 10x",
              "These come from docs/measurements.md: measured at 1x and on tiled 10x frames, with "
              "the database numbers projected. Storage breaks first, at about 1.7 times today's "
              "data on Neon's free tier. Then worker memory, because every job loads the whole "
              "tenant's history. Then the 60-second window a job gets when it starts late in a "
              "tick. The daily sweep itself does not break: about 30 seconds projected at 10x.")
    items = [("1  Neon Free storage, at about 1.7x",
              "584.4 MB at 1x against a 1 GB limit; about 5.7 GB at 10x (projected).",
              "Prune on a schedule, drop the demo feed buffer, then a paid tier with raw_records "
              "and events partitioned by month."),
             ("2  Worker memory, at about 2.9x hearthline",
              "Each job loads the tenant's whole history: 5.0 KB per event. Vercel Hobby's 2 GB "
              "is reached at about 400,000 events; 10x hearthline needs about 6,900 MiB.",
              "Read events straight into columns, only rows visible at T, and scope each rule "
              "check to its one account."),
             ("3  The 60-second late-job window",
              "A tick has a 240 s budget inside a 300 s function, so a job that starts late gets "
              "at most 60 s. An LLM brief must fit in that.",
              "Pass the tick's deadline down and don't start a job whose recent p95 exceeds the "
              "time left. (The brief's evidence query already went from 3.75 s to 1 to 2 ms.)")]
    for i, (head, what, fix) in enumerate(items):
        y = Inches(1.5) + i * Inches(1.8)
        text(s, LEFT, y, Inches(12.0), Inches(1.65), [
            (head, 17, SERIF, RED, False), lead("Limit: ", what, 14),
            lead("Fix: ", fix, 14)], fill=CARD)
    text(s, LEFT, Inches(6.9), Inches(12.0), Inches(0.45), [
        body("The daily sweep does not break: about 30 s for hearthline at 10x, projected from "
             "3.0 s at 1x.", 13, MUTED)])


def s10(prs):
    s = slide(prs, "What I would build next, and what I left out",
              "Next is a real pilot, because the backtest can't tell us whether CSMs act on the "
              "list. After that, the 10x fixes and auto-send for one play at a time once CSMs "
              "send the drafts unedited most of the time. What I left out, I left out on purpose: "
              "fuzzy matching in particular, because a wrong match can put one customer's ticket "
              "into another customer's email.")
    text(s, LEFT, Inches(1.55), Inches(5.8), Inches(5.4), [
        ("Next", 18, SERIF, TEAL, False),
        lead("A six-week pilot ", "with three CSMs and a manager, measured on cases handled on "
             "time and surprise churn: lost ARR that never appeared on its owner's list.", 14),
        lead("The 10x fixes: ", "monthly partitions, columnar loads, a deadline-aware tick.", 14),
        lead("Auto-send, one play at a time, ", "once CSMs send its drafts unedited about 80% of "
             "the time for a month.", 14),
        lead("Which plays work: ", "exposure and outcomes are already logged, so an uplift "
             "model becomes possible.", 14),
        lead("Scheduled retraining ", "behind the existing promotion gates.", 14)])
    x = LEFT + Inches(6.3)
    text(s, x, Inches(1.55), Inches(5.6), Inches(3.4), [
        ("Left out on purpose", 18, SERIF, RED, False),
        lead("Real vendor integrations and real sending. ", "A simulated feed uses the same "
             "ingestion entrypoint as the webhooks.", 14),
        lead("Reading Gmail bodies. ", "Metadata only.", 14),
        lead("Fuzzy identity matching. ", "A missing match is visible and fixable; a wrong one "
             "leaks one customer's data into another's email.", 14),
        lead("An LLM judge, a chat assistant, SSO, ", "and more than one currency per tenant.",
             14)], fill=CARD)


SLIDES = (s1, s2, s3, s4, s5, s6, s7, s8, s9, s10)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n", 1)[0])
    parser.add_argument("--out", type=Path, default=ROOT / "docs" / "deck" / "radar.pptx")
    parser.add_argument("--dump", action="store_true", help="print every slide's text")
    args = parser.parse_args(argv)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs = deck()
    for build in SLIDES:
        build(prs)
    prs.save(args.out)
    problems, dump = check(args.out)
    if args.dump:
        print("\n".join(dump))
    print(f"{args.out}: {len(SLIDES)} slides;", "problems:", problems or "none")
    if problems:
        sys.exit(1)


if __name__ == "__main__":
    main()
